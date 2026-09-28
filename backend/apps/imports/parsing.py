"""Reading the uploaded file: CSV (any common delimiter and encoding) or XLSX (first sheet with data).

The result is the header row (column titles, made unique) and the data rows as `{title: text}` with every
cell as trimmed text: XLSX dates become ISO dates ("2026-10-05"), whole numbers lose the ".0" and decimals
keep two places with a dot (so "1234.5" is never read as thousands). Empty rows are skipped; the row
number is the spreadsheet row (the header is row 1 when it is the first line).
"""

import csv
import io
import re
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal

from apps.core.errors import DomainError
from apps.imports.catalog import MAX_COLUMNS, MAX_FILE_BYTES, MAX_ROWS

ENCODINGS = ("utf-8-sig", "utf-8", "cp1252", "latin-1")
DELIMITERS = [",", ";", "\t", "|"]
XLSX_MAGIC = b"PK\x03\x04"
XLS_MAGIC = b"\xd0\xcf\x11\xe0"  # legacy Excel 97-2003 (BIFF): not supported


class FileError(DomainError):
    code = "invalid_file"


@dataclass
class ParsedFile:
    file_format: str  # csv | xlsx
    headers: list[str]
    rows: list[tuple[int, dict[str, str]]]  # (spreadsheet row number, {title: text})
    sheet_name: str = ""
    delimiter: str = ""
    encoding: str = ""


def _clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, datetime):
        if value.time() == time(0, 0):
            return value.date().isoformat()
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float | Decimal):
        number = Decimal(str(value))
        if number == number.to_integral_value():
            return str(int(number))
        return format(number.quantize(Decimal("0.01")), "f")
    return " ".join(str(value).replace(" ", " ").split())


def _unique_headers(raw: list[str]) -> list[str]:
    """Titles as written (trimmed); empty ones become "Columna N"; repeated ones get " (2)"."""
    headers, seen = [], {}
    for index, title in enumerate(raw, start=1):
        title = _clean(title) or f"Columna {index}"
        count = seen.get(title.lower(), 0) + 1
        seen[title.lower()] = count
        headers.append(title if count == 1 else f"{title} ({count})")
    return headers


def _build(matrix, file_format: str, **extra) -> ParsedFile:
    """First non-empty row = header; the rest = data (empty rows skipped, trailing empty columns dropped)."""
    header_index, header = None, None
    rows: list[tuple[int, dict[str, str]]] = []
    headers: list[str] = []
    for number, raw_row in enumerate(matrix, start=1):
        cells = [_clean(value) for value in raw_row]
        if header is None:
            if not any(cells):
                continue
            while cells and not cells[-1]:
                cells.pop()
            if len(cells) > MAX_COLUMNS:
                raise FileError(
                    f"El archivo tiene {len(cells)} columnas; el máximo es {MAX_COLUMNS}",
                    code="too_many_columns",
                )
            header_index, header = number, cells
            headers = _unique_headers(cells)
            continue
        if not any(cells):
            continue
        if len(rows) >= MAX_ROWS:
            raise FileError(
                f"El archivo tiene más de {MAX_ROWS} filas: divídelo en varios archivos", code="too_many_rows"
            )
        rows.append(
            (number, {title: (cells[i] if i < len(cells) else "") for i, title in enumerate(headers)})
        )
    if header is None:
        raise FileError("El archivo está vacío", code="empty_file")
    if not rows:
        raise FileError("El archivo solo tiene la fila de títulos, sin datos", code="no_rows")
    del header_index
    return ParsedFile(file_format=file_format, headers=headers, rows=rows, **extra)


def _decode(content: bytes) -> tuple[str, str]:
    for encoding in ENCODINGS:
        try:
            return content.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    raise FileError(
        "No se pudo leer el texto del archivo (codificación desconocida)", code="invalid_encoding"
    )


def _delimiter(text: str) -> str:
    sample = "\n".join(text.splitlines()[:20])
    try:
        return csv.Sniffer().sniff(sample, delimiters="".join(DELIMITERS)).delimiter
    except csv.Error:
        first = next((line for line in text.splitlines() if line.strip()), "")
        counts = {delimiter: first.count(delimiter) for delimiter in DELIMITERS}
        best = max(counts, key=counts.get)
        return best if counts[best] else ","


def read_csv(content: bytes) -> ParsedFile:
    text, encoding = _decode(content)
    delimiter = _delimiter(text)
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    try:
        return _build(reader, "csv", delimiter=delimiter, encoding=encoding)
    except csv.Error as exc:
        raise FileError(f"El CSV está mal formado: {exc}", code="invalid_csv") from None


def read_xlsx(content: bytes, sheet: str = "") -> ParsedFile:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - any openpyxl/zip error means "not a readable workbook"
        raise FileError("No se pudo abrir el archivo de Excel (.xlsx)", code="invalid_xlsx") from None
    try:
        sheets = workbook.worksheets
        if sheet:
            chosen = [ws for ws in sheets if ws.title.strip().lower() == sheet.strip().lower()]
            if not chosen:
                raise FileError(f"El libro no tiene una hoja llamada «{sheet}»", code="sheet_not_found")
            candidates = chosen
        else:
            candidates = sheets
        last_error = None
        for worksheet in candidates:
            try:
                return _build(worksheet.iter_rows(values_only=True), "xlsx", sheet_name=worksheet.title)
            except FileError as exc:
                if exc.code not in ("empty_file", "no_rows"):
                    raise
                last_error = exc
        raise last_error or FileError("El archivo está vacío", code="empty_file")
    finally:
        workbook.close()


def read_file(upload, *, sheet: str = "") -> ParsedFile:
    """`upload`: a Django UploadedFile. Raises FileError (400) with a clear message."""
    size = getattr(upload, "size", None) or 0
    if size > MAX_FILE_BYTES:
        raise FileError(f"El archivo pesa más de {MAX_FILE_BYTES // (1024 * 1024)} MB", code="file_too_large")
    content = upload.read(MAX_FILE_BYTES + 1)
    if len(content) > MAX_FILE_BYTES:
        raise FileError(f"El archivo pesa más de {MAX_FILE_BYTES // (1024 * 1024)} MB", code="file_too_large")
    if not content.strip():
        raise FileError("El archivo está vacío", code="empty_file")
    name = (getattr(upload, "name", "") or "").lower()
    if content.startswith(XLS_MAGIC) or name.endswith(".xls"):
        raise FileError(
            "Los archivos .xls (Excel 97-2003) no se pueden leer: guárdalo como .xlsx o .csv",
            code="unsupported_format",
        )
    if content.startswith(XLSX_MAGIC):
        return read_xlsx(content, sheet)
    if name.endswith((".xlsx", ".xlsm")):
        raise FileError("No se pudo abrir el archivo de Excel (.xlsx)", code="invalid_xlsx")
    if b"\x00" in content[:4096]:
        raise FileError("El archivo no es un CSV ni un Excel (.xlsx)", code="unsupported_format")
    return read_csv(content)


def safe_filename(name: str) -> str:
    base = (name or "archivo").replace("\\", "/").rsplit("/", 1)[-1]
    return re.sub(r"[^\w.\- ()áéíóúñÁÉÍÓÚÑ]+", "_", base)[:200] or "archivo"
