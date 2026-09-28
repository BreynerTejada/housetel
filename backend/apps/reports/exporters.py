"""CSV, XLSX and PDF exports of a report, rendered from the same ``ReportResult`` as the JSON.

- **CSV**: ``;``-separated, UTF-8 with BOM (Excel in Spanish opens it with accents and columns right). Numbers
  use the report language's decimal separator (es → ``72,5``), COP money is whole pesos (``320000``) and dates
  are ISO (``2026-09-30``), which Excel reads in any locale. One table: header + rows (+ ``Total`` row);
  several tables: each one preceded by its title and separated by a blank line.
- **XLSX** (openpyxl): a summary sheet (property, range, KPIs, definitions) and one sheet per table with real
  numbers and dates, number formats, bold header and totals, frozen header and fitted column widths.
- **PDF** (reportlab): Housetel/property header on every page, KPIs, the tables (header repeated on each page)
  and the definitions; landscape when a table is wide.
"""

from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal
from xml.sax.saxutils import escape

from django.http import HttpResponse
from django.utils import timezone

from apps.core.money import D
from apps.reports.builders.common import day_text, month_text, prev_title
from apps.reports.labels import country_label, report_title, tr
from apps.reports.output import (
    BOOLEAN,
    CODE,
    COUNTRY,
    DATETIME,
    MONEY,
    MONTH,
    NUMBER,
    PERCENT,
    STATUS,
    Column,
    ReportResult,
    Table,
)

CONTENT_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}
PDF_MAX_ROWS = 2000
BRAND = "#B4583B"
INK = "#1F1C19"
MUTED = "#6E675E"
HAIRLINE = "#E7E2DA"
HEADER_FILL = "#F3F0EB"


def export_filename(report_id: str, p, fmt: str) -> str:
    parts = [report_id, p.prop.slug]
    if p.start is not None:
        parts += [p.start.isoformat(), p.end.isoformat()]
    else:
        parts.append(p.business_date.isoformat())
    return f"{'_'.join(parts)}.{fmt}"


def export_report(report, p, result: ReportResult, fmt: str) -> HttpResponse:
    if fmt == "csv":
        content = render_csv(result, p)
    elif fmt == "xlsx":
        content = render_xlsx(report, p, result)
    else:
        content = render_pdf(report, p, result)
    response = HttpResponse(content, content_type=CONTENT_TYPES[fmt])
    response["Content-Disposition"] = f'attachment; filename="{export_filename(report.id, p, fmt)}"'
    response["Cache-Control"] = "no-store"
    return response


# --- value formatting ---------------------------------------------------------------------------------------


def _decimal_text(value: Decimal, digits: int, lang: str, grouping: bool = False) -> str:
    quantum = Decimal(1).scaleb(-digits) if digits else Decimal(1)
    number = D(value).quantize(quantum, rounding=ROUND_HALF_UP)
    text = f"{number:,.{digits}f}" if grouping else f"{number:.{digits}f}"
    if lang == "en":
        return text
    return text.replace(",", "\0").replace(".", ",").replace("\0", ".")


def _label(column: Column, value) -> str:
    if column.labels and value in column.labels:
        return column.labels[value]
    return "" if value is None else str(value)


def csv_value(value, column: Column, lang: str, currency: str) -> str:
    if value is None or value == "":
        return ""
    kind = column.type
    if kind == MONEY:
        return _decimal_text(value, 0 if currency == "COP" else 2, lang)
    if kind == PERCENT:
        return _decimal_text(value, 1, lang)
    if kind == NUMBER:
        if isinstance(value, Decimal) and value != value.to_integral_value():
            return _decimal_text(value, 1, lang)
        return str(int(value)) if isinstance(value, Decimal) else str(value)
    if kind == DATETIME and isinstance(value, datetime):
        return timezone.localtime(value).strftime("%Y-%m-%d %H:%M")
    if kind == MONTH and isinstance(value, date):
        return value.strftime("%Y-%m")
    if isinstance(value, date):
        return value.isoformat()
    if kind in (STATUS, CODE):
        return _label(column, value)
    if kind == COUNTRY:
        return country_label(value, lang)
    if kind == BOOLEAN:
        return tr("yes", lang) if value else tr("no", lang)
    return str(value)


def display_value(value, column: Column, lang: str, currency: str) -> str:
    """Human text for the PDF (money like the app: ``$ 320.000``)."""
    if value is None or value == "":
        return ""
    kind = column.type
    if kind == MONEY:
        digits = 0 if currency == "COP" else 2
        amount = D(value)
        text = _decimal_text(abs(amount), digits, "es", grouping=True)
        return f"{'-' if amount < 0 else ''}$ {text}"
    if kind == PERCENT:
        text = _decimal_text(value, 1, lang)
        return f"{text} %" if lang != "en" else f"{text}%"
    if kind == NUMBER:
        if isinstance(value, Decimal) and value != value.to_integral_value():
            return _decimal_text(value, 1, lang, grouping=True)
        return _decimal_text(D(value), 0, lang, grouping=True)
    if kind == DATETIME and isinstance(value, datetime):
        local = timezone.localtime(value)
        return f"{day_text(local.date(), lang)} {local:%H:%M}"
    if kind == MONTH and isinstance(value, date):
        return month_text(value, lang)
    if isinstance(value, date):
        return day_text(value, lang)
    return csv_value(value, column, lang, currency)


def _first_cell_total(table: Table, lang: str) -> dict:
    """The totals row with "Total" in its first cell, unless a text column already carries the label."""
    totals = dict(table.totals or {})
    first = table.columns[0].key
    labelled = any(
        column.type == "text" and totals.get(column.key) not in (None, "") for column in table.columns
    )
    if totals.get(first) in (None, "") and not labelled:
        totals[first] = tr("total_row", lang)
    return totals


def _kpi_text(kpi, lang: str, currency: str) -> str:
    column = Column(kpi.key, kpi.label, kpi.type)
    return display_value(kpi.value, column, lang, currency) if kpi.type != "text" else str(kpi.value or "")


# --- CSV ----------------------------------------------------------------------------------------------------


def render_csv(result: ReportResult, p) -> bytes:
    buffer = io.StringIO()
    buffer.write("\ufeff")
    writer = csv.writer(buffer, delimiter=";", lineterminator="\r\n")
    several = len(result.tables) > 1
    for index, table in enumerate(result.tables):
        if several:
            if index:
                writer.writerow([])
            writer.writerow([table.title])
        writer.writerow([column.label for column in table.columns])
        for row in table.rows:
            writer.writerow([csv_value(row.get(c.key), c, p.lang, p.currency) for c in table.columns])
        if table.totals:
            totals = _first_cell_total(table, p.lang)
            first = table.columns[0].key
            writer.writerow(
                [
                    totals.get(c.key)
                    if c.key == first
                    else csv_value(totals.get(c.key), c, p.lang, p.currency)
                    for c in table.columns
                ]
            )
    return buffer.getvalue().encode("utf-8")


# --- XLSX ---------------------------------------------------------------------------------------------------

_SHEET_INVALID = re.compile(r"[\[\]:*?/\\]")


def _sheet_name(title: str, used: set) -> str:
    base = _SHEET_INVALID.sub(" ", title).strip()[:31] or "Datos"
    name, counter = base, 2
    while name.lower() in used:
        suffix = f" ({counter})"
        name = f"{base[: 31 - len(suffix)]}{suffix}"
        counter += 1
    used.add(name.lower())
    return name


def _xlsx_cell(value, column: Column, lang: str, currency: str):
    """(value, number_format) with real numbers and dates for Excel."""
    if value is None or value == "":
        return None, None
    kind = column.type
    money_format = '"$" #,##0' if currency == "COP" else '"$" #,##0.00'
    if kind == MONEY:
        return float(D(value)), money_format
    if kind == PERCENT:
        return float(D(value) / 100), "0.0%"
    if kind == NUMBER:
        if isinstance(value, Decimal) and value != value.to_integral_value():
            return float(value), "#,##0.0"
        return int(value), "#,##0"
    if kind == DATETIME and isinstance(value, datetime):
        return timezone.localtime(value).replace(tzinfo=None), "yyyy-mm-dd hh:mm"
    if kind == MONTH and isinstance(value, date):
        return value, "yyyy-mm"
    if isinstance(value, date):
        return value, "yyyy-mm-dd"
    return csv_value(value, column, lang, currency), None


def render_xlsx(report, p, result: ReportResult) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    lang, currency = p.lang, p.currency
    bold = Font(bold=True, color=INK[1:])
    title_font = Font(bold=True, size=14, color=INK[1:])
    brand_font = Font(bold=True, size=11, color=BRAND[1:])
    muted = Font(color=MUTED[1:], size=10)
    fill = PatternFill("solid", fgColor=HEADER_FILL[1:])
    thin = Side(style="thin", color=HAIRLINE[1:])
    medium = Side(style="medium", color=INK[1:])

    workbook = Workbook()
    used: set = set()
    summary = workbook.active
    summary.title = _sheet_name("Resumen" if lang != "en" else "Summary", used)
    summary["A1"] = f"Housetel · {p.prop.name}"
    summary["A1"].font = brand_font
    summary["A2"] = report_title(report.id, lang)
    summary["A2"].font = title_font
    row = 3
    for note in result.notes[:2]:
        summary.cell(row=row, column=1, value=note).font = muted
        row += 1
    generated = timezone.localtime()
    summary.cell(
        row=row,
        column=1,
        value=f"{'Generado' if lang != 'en' else 'Generated'}: {generated:%Y-%m-%d %H:%M}",
    ).font = muted
    row += 2
    if result.summary:
        compared = p.compared
        headers = [("Indicador" if lang != "en" else "Indicator"), ("Valor" if lang != "en" else "Value")]
        if compared:
            headers += [prev_title(p), "Variación" if lang != "en" else "Change"]
        for col, header in enumerate(headers, start=1):
            cell = summary.cell(row=row, column=col, value=header)
            cell.font, cell.fill, cell.border = bold, fill, Border(bottom=thin)
        row += 1
        for kpi in result.summary:
            column = Column(kpi.key, kpi.label, kpi.type)
            summary.cell(row=row, column=1, value=kpi.label)
            value, number_format = (
                _xlsx_cell(kpi.value, column, lang, currency) if kpi.type != "text" else (kpi.value, None)
            )
            cell = summary.cell(row=row, column=2, value=value)
            if number_format:
                cell.number_format = number_format
            if compared:
                prev_value, prev_format = (
                    _xlsx_cell(kpi.previous, column, lang, currency)
                    if kpi.type != "text"
                    else (kpi.previous, None)
                )
                prev_cell = summary.cell(row=row, column=3, value=prev_value)
                if prev_format:
                    prev_cell.number_format = prev_format
                if kpi.type == PERCENT and kpi.value is not None and kpi.previous is not None:
                    change_cell = summary.cell(row=row, column=4, value=float(D(kpi.value) - D(kpi.previous)))
                    change_cell.number_format = '+0.0" pp";-0.0" pp";0.0" pp"'
                elif (
                    kpi.type != "text"
                    and kpi.value is not None
                    and kpi.previous not in (None, 0)
                    and D(kpi.previous)
                ):
                    change = (D(kpi.value) - D(kpi.previous)) / abs(D(kpi.previous))
                    change_cell = summary.cell(row=row, column=4, value=float(change))
                    change_cell.number_format = "+0.0%;-0.0%;0.0%"
            row += 1
        row += 1
    for note in result.notes[2:]:
        cell = summary.cell(row=row, column=1, value=note)
        cell.font = muted
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        summary.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        summary.row_dimensions[row].height = 15 * max(1, len(note) // 110 + 1)
        row += 1
    summary.column_dimensions["A"].width = 44
    for letter in ("B", "C", "D"):
        summary.column_dimensions[letter].width = 22

    for table in result.tables:
        sheet = workbook.create_sheet(_sheet_name(table.title, used))
        sheet["A1"] = table.title
        sheet["A1"].font = title_font
        header_row = 3
        widths = [len(column.label) for column in table.columns]
        for col, column in enumerate(table.columns, start=1):
            cell = sheet.cell(row=header_row, column=col, value=column.label)
            cell.font, cell.fill = bold, fill
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(wrap_text=True, vertical="bottom")
        current = header_row + 1
        rows = list(table.rows)
        for data in rows:
            for col, column in enumerate(table.columns, start=1):
                value, number_format = _xlsx_cell(data.get(column.key), column, lang, currency)
                cell = sheet.cell(row=current, column=col, value=value)
                if number_format:
                    cell.number_format = number_format
                text = display_value(data.get(column.key), column, lang, currency)
                widths[col - 1] = max(widths[col - 1], len(text))
            current += 1
        if table.totals:
            totals = _first_cell_total(table, lang)
            for col, column in enumerate(table.columns, start=1):
                raw = totals.get(column.key)
                if col == 1:
                    value, number_format = raw, None
                else:
                    value, number_format = _xlsx_cell(raw, column, lang, currency)
                cell = sheet.cell(row=current, column=col, value=value)
                cell.font = bold
                cell.border = Border(top=medium)
                if number_format:
                    cell.number_format = number_format
        for col, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(col)].width = min(max(width + 2, 9), 48)
        sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
        if rows:
            sheet.auto_filter.ref = (
                f"A{header_row}:{get_column_letter(len(table.columns))}{header_row + len(rows)}"
            )

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# --- PDF ----------------------------------------------------------------------------------------------------

_PDF_REPLACEMENTS = {"≤": "<=", "≥": ">=", "−": "-", "→": "->", " ": " ", " ": " ", "…": "..."}


def _pdf_text(text) -> str:
    text = "" if text is None else str(text)
    for old, new in _PDF_REPLACEMENTS.items():
        text = text.replace(old, new)
    return escape(text).replace("$ ", "$&nbsp;")


def render_pdf(report, p, result: ReportResult) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, TableStyle
    from reportlab.platypus import Table as PdfTable

    lang, currency = p.lang, p.currency
    wide = max((len(table.columns) for table in result.tables), default=0) > 7
    pagesize = landscape(A4) if wide else A4
    margin = 14 * mm
    width = pagesize[0] - 2 * margin
    brand, ink, muted, hairline = (colors.HexColor(c) for c in (BRAND, INK, MUTED, HAIRLINE))
    header_fill = colors.HexColor(HEADER_FILL)
    zebra = colors.HexColor("#FAF8F5")

    base = ParagraphStyle("base", fontName="Helvetica", fontSize=8.5, leading=11, textColor=ink)
    small = ParagraphStyle("small", parent=base, fontSize=7.5, leading=9.5, textColor=muted)
    title = ParagraphStyle(
        "title", parent=base, fontName="Helvetica-Bold", fontSize=17, leading=21, spaceAfter=2
    )
    heading = ParagraphStyle(
        "heading",
        parent=base,
        fontName="Helvetica-Bold",
        fontSize=10.5,
        leading=14,
        spaceBefore=10,
        spaceAfter=4,
    )
    cell = ParagraphStyle("cell", parent=base, fontSize=7.3, leading=9)
    cell_right = ParagraphStyle("cell_right", parent=cell, alignment=TA_RIGHT)
    cell_head = ParagraphStyle("cell_head", parent=cell, fontName="Helvetica-Bold", textColor=ink)
    cell_head_right = ParagraphStyle("cell_head_right", parent=cell_head, alignment=TA_RIGHT)
    kpi_label = ParagraphStyle("kpi_label", parent=small, fontSize=7.5, textColor=muted)
    kpi_value = ParagraphStyle("kpi_value", parent=base, fontName="Helvetica-Bold", fontSize=12.5, leading=15)

    report_name = report_title(report.id, lang)
    generated = timezone.localtime()

    def decorate(canvas, doc):
        canvas.saveState()
        top = pagesize[1] - margin + 4 * mm
        canvas.setFillColor(brand)
        canvas.setFont("Helvetica-Bold", 9.5)
        canvas.drawString(margin, top, "Housetel")
        canvas.setFillColor(ink)
        canvas.setFont("Helvetica", 9)
        canvas.drawString(margin + 44, top, str(p.prop.name))
        canvas.setFillColor(muted)
        canvas.drawRightString(pagesize[0] - margin, top, report_name)
        canvas.setStrokeColor(brand)
        canvas.setLineWidth(0.8)
        canvas.line(margin, top - 2.5 * mm, pagesize[0] - margin, top - 2.5 * mm)
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(muted)
        stamp = f"{'Generado' if lang != 'en' else 'Generated'} {generated:%Y-%m-%d %H:%M}"
        canvas.drawString(margin, margin - 6 * mm, stamp)
        page = f"{'Página' if lang != 'en' else 'Page'} {doc.page}"
        canvas.drawRightString(pagesize[0] - margin, margin - 6 * mm, page)
        canvas.restoreState()

    story = [Paragraph(_pdf_text(report_name), title)]
    for note in result.notes[:2]:
        story.append(Paragraph(_pdf_text(note), small))
    story.append(Spacer(1, 5 * mm))

    kpis = [k for k in result.summary]
    if kpis:
        per_row = 4 if not wide else 6
        cells, grid = [], []
        for kpi in kpis:
            block = [
                Paragraph(_pdf_text(kpi.label), kpi_label),
                Paragraph(_pdf_text(_kpi_text(kpi, lang, currency)), kpi_value),
            ]
            if p.compared and kpi.previous is not None and kpi.type != "text":
                prev_text = display_value(kpi.previous, Column(kpi.key, kpi.label, kpi.type), lang, currency)
                block.append(Paragraph(_pdf_text(f"{prev_title(p)}: {prev_text}"), small))
            cells.append(block)
        for start in range(0, len(cells), per_row):
            chunk = cells[start : start + per_row]
            grid.append(chunk + [""] * (per_row - len(chunk)))
        kpi_table = PdfTable(grid, colWidths=[width / per_row] * per_row)
        kpi_table.setStyle(
            TableStyle(
                [
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LINEABOVE", (0, 0), (-1, 0), 0.6, hairline),
                    ("LINEBELOW", (0, -1), (-1, -1), 0.6, hairline),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                    ("LEFTPADDING", (0, 0), (-1, -1), 2),
                ]
            )
        )
        story += [kpi_table, Spacer(1, 3 * mm)]

    numeric = (MONEY, PERCENT, NUMBER)
    for table in result.tables:
        rows = table.rows[:PDF_MAX_ROWS]
        header = [
            Paragraph(_pdf_text(c.label), cell_head_right if c.type in numeric else cell_head)
            for c in table.columns
        ]
        body = []
        texts = []
        for data in rows:
            values = [display_value(data.get(c.key), c, lang, currency) for c in table.columns]
            texts.append(values)
            body.append(
                [
                    Paragraph(_pdf_text(v), cell_right if c.type in numeric else cell)
                    for v, c in zip(values, table.columns, strict=True)
                ]
            )
        if table.totals:
            totals = _first_cell_total(table, lang)
            values = [
                str(totals.get(c.key) or "")
                if i == 0
                else display_value(totals.get(c.key), c, lang, currency)
                for i, c in enumerate(table.columns)
            ]
            texts.append(values)
            body.append(
                [
                    Paragraph(f"<b>{_pdf_text(v)}</b>", cell_right if c.type in numeric else cell)
                    for v, c in zip(values, table.columns, strict=True)
                ]
            )
        # Column widths proportional to their longest text (the header counts less, but never less than its
        # longest word, so a header wraps between words and never inside one), fitted to the page.
        weights = []
        for index, column in enumerate(table.columns):
            longest_word = max((len(word) for word in column.label.split()), default=0) * 1.15
            longest = max(
                [len(r[index]) for r in texts] + [min(len(column.label), 16) * 0.75, longest_word, 6]
            )
            weights.append(min(longest, 42))
        total_weight = sum(weights) or 1
        col_widths = [width * w / total_weight for w in weights]
        pdf_table = PdfTable([header, *body], colWidths=col_widths, repeatRows=1)
        style = [
            ("BACKGROUND", (0, 0), (-1, 0), header_fill),
            ("LINEBELOW", (0, 0), (-1, 0), 0.6, hairline),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 2.2),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2.2),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ]
        for index in range(1, len(body) + 1):
            if index % 2 == 0:
                style.append(("BACKGROUND", (0, index), (-1, index), zebra))
        if table.totals:
            style.append(("LINEABOVE", (0, len(body)), (-1, len(body)), 0.8, ink))
        pdf_table.setStyle(TableStyle(style))
        block = [Paragraph(_pdf_text(table.title), heading)]
        if not rows:
            block.append(
                Paragraph(
                    _pdf_text("Sin datos en este rango." if lang != "en" else "No data in this range."), small
                )
            )
            story.append(KeepTogether(block))
            continue
        story.append(KeepTogether([block[0], pdf_table]) if len(body) < 25 else block[0])
        if len(body) >= 25:
            story.append(pdf_table)
        if len(table.rows) > PDF_MAX_ROWS:
            story.append(
                Paragraph(
                    _pdf_text(
                        f"Se muestran {PDF_MAX_ROWS} de {len(table.rows)} filas; exporta a Excel para verlas "
                        "todas."
                        if lang != "en"
                        else f"Showing {PDF_MAX_ROWS} of {len(table.rows)} rows; "
                        "export to Excel to see them all."
                    ),
                    small,
                )
            )

    if len(result.notes) > 2:
        story.append(
            Paragraph(_pdf_text("Cómo se calcula" if lang != "en" else "How it is calculated"), heading)
        )
        for note in result.notes[2:]:
            story.append(Paragraph(_pdf_text(note), small))
            story.append(Spacer(1, 1.5 * mm))

    buffer = io.BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=pagesize,
        leftMargin=margin,
        rightMargin=margin,
        topMargin=margin + 6 * mm,
        bottomMargin=margin,
        title=f"{report_name} · {p.prop.name}",
        author="Housetel",
    )
    document.build(story, onFirstPage=decorate, onLaterPages=decorate)
    return buffer.getvalue()
